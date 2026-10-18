import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tariffapp import db  # noqa: E402


def make_indigo_file(path, rows):
    """Minimal IndiGo rate list: a GCR+LEAN sheet with the usual headers.
    rows: (org, dest, type, min, n, q45, q100, applicable_flights, sector_surcharge, flight_surcharge)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GCR+LEAN+SCR+DGR"
    ws.append(["IndiGo Cargo domestic rates"])
    ws.append(["Region", "ORG", "DEST", "Type", "MIN", "N", "Q45", "Q100", "Q300", "Q500", "Q1000",
               "Applicable Flights", "Applicable Commodity", "Sector Surcharge", "Flight Surcharge",
               "Commodity Surcharge"])
    for o, d, t, mn, n, q45, q100, cond, sec, flt in rows:
        ws.append(["West", o, d, t, mn, n, q45, q100, None, None, None, cond, None, sec, flt, None])
    fn = wb.create_sheet("Footnotes")
    fn.append(["IndiGo Cargo Footnotes"])
    fn.append([1, "Fuel surcharge of INR 4/kg applicable PAN India"])
    wb.save(path)


def make_airindia_file(path, rows):
    """rows: (org, dest, category, min, n, q45, q100, flight_number_text, surcharge, remarks)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tariff Sheet"
    ws.append(["Air India domestic tariff"])
    ws.append(["Region", "Origin Code", "Destination Code", "Sector", "Applicable Flight Number",
               "Prime/GCR/ Lean/SCR", "Applicable Commodity", "MIN", "N", "Q45", "Q100", "Q300", "Q500",
               "Q1000", "Surcharge", "Seasonal Surcharge", "Remarks"])
    for o, d, cat, mn, n, q45, q100, flt, sur, rem in rows:
        ws.append(["North", o, d, f"{o}-{d}", flt, cat, None, mn, n, q45, q100, None, None, None, sur, None, rem])
    wb.save(path)


def make_schedule_file(path, flights):
    """flights: (origin, dest, airline, flight_no, dep_time, days_string 'MTWTFSS' with '.' for off)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw Data"
    ws.append(["Origin", "Destination", "OD Pair", "Airline", "Flight No", "Dep Time",
               "Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun", "Weekly Freq"])
    for o, d, a, f, t, days in flights:
        ws.append([o, d, f"{o}-{d}", a, f, t, *[0 if c == "." else 1 for c in days], None])
    ws.append(["AMD", "IXZ", "AMD-IXZ", "NO_SERVICE", None, None, 0, 0, 0, 0, 0, 0, 0, 0])
    lists = wb.create_sheet("Lists")
    lists.append(["Code", "Name", None, "ATC Codes", None, "Destination Codes"])
    for code, name, hub in [("AMD", "Ahmedabad", True), ("DEL", "Delhi", True), ("BOM", "Mumbai", True),
                            ("COK", "Kochi", False), ("TVM", "Thiruvananthapuram", False)]:
        lists.append([code, name, None, code if hub else None, None, code])
    addon = wb.create_sheet("Add-on Rates")
    addon.append(["Airline code (as in Raw Data)", "Airline", "Add-on ₹/kg", "Includes"])
    addon.append(["AIRINDIA", "Air India", 12, "handling"])
    addon.append(["INDIGO", "IndiGo", 18, "handling"])
    addon.append(["AKASA", "Akasa Air", None, "not set"])
    addon.append([])
    addon.append(["Edit the yellow values when your add-on costs change"])
    ns = wb.create_sheet("Network Surcharges")
    ns.append(["Airline", "Applies To", "Station / Route", "Rs per kg", "Valid From", "Valid To", "Reason",
               "Active (Y/N)", "Dashboard Label"])
    ns.append(["IndiGo", "All", None, 4, "2026-04-02", None, "Fuel surcharge", "Y", "Fuel"])
    ns.append(["Air India", "All", None, 5, "2026-04-03", None, "SS surcharge", "Y", "SS"])
    ns.append(["Air India", "Destination", "BOM", 6, "2026-08-20", None, "To BOM", "Y", "To-BOM"])
    ns.append(["Air India", "Route", "DEL-COK", 5, "2026-08-07", "2026-08-31", "Onam", "Y", "Onam"])
    ns.append(["IndiGo", "Origin", "AMD", 99, None, None, "switched off", "N", "Off"])
    wb.save(path)


@pytest.fixture
def db_path(tmp_path, monkeypatch):
    path = str(tmp_path / "t.sqlite")
    monkeypatch.setenv("TARIFFAPP_DB", path)
    monkeypatch.setenv("TARIFFAPP_DATA", str(tmp_path / "data"))  # never touch the repo's data folder
    return path


@pytest.fixture
def conn(db_path):
    c = db.connect(db_path)
    db.init_schema(c)
    yield c
    c.close()


@pytest.fixture
def loaded(conn, tmp_path):
    """A database with one schedule + two IndiGo cycles + one Air India cycle."""
    from tariffapp.etl.loader import load_rate_file
    from tariffapp.importers import import_schedule_workbook

    sched = tmp_path / "schedule.xlsx"
    make_schedule_file(sched, [
        ("AMD", "DEL", "INDIGO", "6E396", "05:00", "MTWTFSS"),
        ("AMD", "DEL", "INDIGO", "6E100", "14:00", "MTW...."),
        ("AMD", "DEL", "INDIGO", "6E5245", "23:30", "......S"),
        ("AMD", "DEL", "INDIGO", "6E777", "0", "MTWTFSS"),
        ("AMD", "DEL", "AKASA", "QP1000", "09:00", "MTWTFSS"),
        ("DEL", "COK", "AIRINDIA", "AI1715", "07:20", "M......"),
        ("DEL", "COK", "AIRINDIA", "AI2458", "12:10", "MTWTFSS"),
        ("DEL", "BOM", "AIRINDIA", "AI2421", "06:00", "MTWTFSS"),
        ("AMD", "TVM", "INDIGO", "6E50", "10:00", "MTWTFSS"),
    ])
    import_schedule_workbook(conn, str(sched))

    ig1 = tmp_path / "7.0 All India Rate List Eff. 01st July 2026.xlsx"
    make_indigo_file(ig1, [
        ("AMD", "DEL", "GCR", 1500, 70, 50, 40, None, None, None),
    ])
    ig2 = tmp_path / "8.1 6E Eff. 16th Aug 2026.xlsx"
    make_indigo_file(ig2, [
        ("AMD", "DEL", "GCR", 1800, 71, 57, 43, "For All Flights departed between 01:01 TO 06:30 HRS", None, None),
        ("AMD", "DEL", "LEAN", 1800, 69, 55, 26, "For All Flights departed between 06:31 TO 01:00 HRS",
         "INR.5/- Per kg", "INR 5/- per kg (PEF) on 6E100"),
        ("AMD", "DEL", "GCR", 1800, 90, 80, 70, "Applicable for 6E 5245", None, None),
        ("AMD", "TRV", "GCR", 1800, 68, None, None, None, None, None),
    ])
    ai = tmp_path / "8.3 AI effective 20-Aug-26.xlsx"
    make_airindia_file(ai, [
        ("DEL", "COK", "GCR", 1500, 69, 57, 48, None, "INR 10 / KG",
         "Increase commercial surcharge from INR 5 to INR 10 per kg Effective 16-Jul to 31-Aug 2026"),
        ("DEL", "COK", "PRIME", 1500, 82, 72, 61, "Applicable Flight \n(AI-0420, AI-1715, IX-1812)", "INR 10 / KG",
         None),
        ("DEL", "BOM", "GCR", 1500, 71, 55, 25, None, "INR 3 / KG",
         "Introduction of commercial surcharge of INR 3/kg Eff. 02-Jun-26 (Only for FRUITS, VEG, PER, CHR)"),
        ("DEL", "BOM", "PRIME", 1500, 71, 55, 45,
         "Applicable Flights 00:00 to 08:00 Hours\n(AI-2421, AI-2429) Excluding AI-0423", None, None),
    ])
    for f in (ig1, ig2, ai):
        r = load_rate_file(conn, str(f))
        assert r["status"] == "ok", r
    return conn
