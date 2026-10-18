"""Akasa, SpiceJet and Star Air tariffs: reading the files and pricing with them."""

import datetime as dt

import openpyxl

from tariffapp import db
from tariffapp.etl.loader import load_rate_file
from tariffapp.linker import TariffEngine


def _akasa(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tariff Rate 01st Oct 2026"
    ws.append(["TARIFF SHEET"])
    ws.append(["Origin", "Destination", "Type", "Minimum", "Flight No", "Normal Rate", "45 + ", "100 +", "REMARK"])
    ws.append(["BLR", "AMD", "GCR", 1400, None, 60, 50, 30, None])
    ws.append(["BLR", "AMD", "Prime", 1400, "QP1902", 67, 65, 45, None])
    ws.append(["CCU", "DEL", "GCR", 1400, None, 70, 60, 40, None])
    wb.create_sheet("OCDC").append(["Outbound"])
    s = wb.create_sheet("Surcharges")
    s.append(["Surcharges"])
    s.append(["NEWSr. No.", "ORIGIN", "DEST", "FLIGHT NO.", "FLT TYPE", "COMMODITY", "From Date", "To Date",
              "Rate P/kg", "REMARK"])
    s.append([1, "BLR", "AMD", "QP 1902 ", None, None, dt.datetime(2026, 6, 1), "UFN", "Rs.6/-"])
    s.append([2, "ALL", "ALL", "All", None, "Mobile & accessories", dt.datetime(2024, 9, 20), "UFN", "Rs.3/-"])
    s.append([3, "CCU", "DEL", None, None, "EXCLUDING QP1804", dt.datetime(2025, 11, 10), "UFN", "Rs.8/-"])
    s.append([4, "BLR", "CCU", "PER / PEF ", None, None, dt.datetime(2025, 7, 1), "UFN", "Rs.5/-"])
    wb.save(path)


def _spicejet(path):
    wb = openpyxl.Workbook()
    p = wb.active
    p.title = "PRIME FLIGHTS"
    p.append(["PRIME FLIGHTS TARIFF"])
    p.append(["EFFECTIVE : 19-Apr-2026"])
    p.append(["Origin ", "Destination", "Flt No", "MIN", "N Rate", "+45", "+100", "+300"])
    p.append(["BOM ", "GAU", "All Flights", 1500, 106, 88, 82, None])
    g = wb.create_sheet("GCR")
    g.append(["GENERAL CARGO RATES"])
    g.append(["Origin ", "Destination", "MIN", "N Rate", "+45", "+100", "+300", "+500", "+1000", "Commodity",
              "Flight", "Surcharge "])
    g.append(["BOM", "GAU", 1500, 102, 85, 67, None, None, None, None, None, None])
    g.append(["BOM", "GAU", 1500, 60, 50, 40, None, None, None, "Fruits & Veg", None, None])
    wb.create_sheet("VIA FLIGHTS RATES").append(["CONNECTING (VIA) FLIGHTS"])
    h = wb.create_sheet("Handling Charges")
    h.append(["Category", "Rate", "Remarks"])
    h.append(["Fuel Surcharge (FSC)", "INR 3 PER KG", "From:  AMD, BLR, BOM, CCU"])
    h.append(["Security Surcharge (SSC)", "INR 1.50 PER KG"])
    wb.save(path)


def _starair(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GCR &PER PUBLISHED RATES"
    ws.append([None, "STARAIR CARGO PUBLISHED TARIFF -EFF 08th June 2026-UFN"])
    ws.append([None, "ORG", "DEST", "COMMODITY", "CATEGORY", "MINIMUM", "NORMAL", "+45", "+100", "+250"])
    ws.append([None, "AMD", "BLR", "GCR & PER", "GEN", 1200, 45, 32, 22, 18])
    f = wb.create_sheet("FOOTNOTE")
    f.append(["Sector Surcharge (SS)", "Rs 1/- per kg on All sectors-UFN"])
    f.append(["Fuel Surcharge (FS)", "Rs 3/- per Kg on All Sectors-UFN"])
    wb.save(path)


def _price(conn, flight, day=dt.date(2026, 10, 14)):
    f = {"Origin": flight[0], "Destination": flight[1], "Airline": flight[2], "FlightNo": flight[3],
         "DepTime": "10:00", **{d: 1 for d in ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]}}
    return TariffEngine(conn, day).price(f)


def test_akasa_prime_flight_and_surcharges(conn, tmp_path):
    p = tmp_path / "Akasa Revised Tariff_Rates_effective_01st_Oct_2026.xlsx"
    _akasa(p)
    r = load_rate_file(conn, str(p))
    assert (r["status"], r["carrier"], str(r["effective_date"])) == ("ok", "Akasa", "2026-10-01")
    assert conn.execute("SELECT TariffCarrier FROM Airlines WHERE Code='AKASA'").fetchone()[0] == "Akasa"
    prime = _price(conn, ("BLR", "AMD", "AKASA", "QP1902"))
    assert (prime.product, prime.rates["100"], prime.surcharge_total) == ("PRIME", 45, 6)   # flight surcharge
    other = _price(conn, ("BLR", "AMD", "AKASA", "QP1900"))
    assert (other.product, other.rates["100"], other.surcharge_total) == ("GCR", 30, 0)
    # 'EXCLUDING QP1804' (typed in the commodity column) is a flight rule, not a commodity
    assert _price(conn, ("CCU", "DEL", "AKASA", "QP1800")).surcharge_total == 8
    assert _price(conn, ("CCU", "DEL", "AKASA", "QP1804")).surcharge_total == 0


def test_spicejet_all_flights_prime_wins_and_surcharges(conn, tmp_path):
    p = tmp_path / "SG Published Tariff 29-Jul-26.xlsx"
    _spicejet(p)
    assert load_rate_file(conn, str(p))["carrier"] == "SpiceJet"
    r = _price(conn, ("BOM", "GAU", "SPICEJET", "SG123"))
    # PRIME 'All Flights' beats the route GCR; the Fruits & Veg rate is commodity-only and ignored
    assert (r.status, r.product, r.rates["100"]) == ("OK", "PRIME", 82)
    assert r.surcharge_summary == "Fuel ₹3 + Security ₹1.5"  # fuel only from the listed origins
    assert _price(conn, ("GAU", "BOM", "SPICEJET", "SG1")).status == "NO RATE"


def test_star_air_date_in_file_and_250_break(conn, tmp_path):
    p = tmp_path / "STAR AIR PUBLISHED CARGO TARIFF W.E.F. 8TH JUNE-UFN.xlsx"
    _starair(p)
    r = load_rate_file(conn, str(p))
    assert (r["carrier"], str(r["effective_date"])) == ("Star Air", "2026-06-08")
    f = _price(conn, ("AMD", "BLR", "STARAIR", "S5101"))
    assert f.rate_at(300) == (18, "250+") and f.surcharge_total == 4
    assert f.freight_charge(260) == 18 * 260


def test_new_carriers_survive_a_fresh_schema(tmp_path):
    c = db.connect(str(tmp_path / "x.sqlite"))
    db.init_schema(c)
    tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert {"Akasa_GCR", "SpiceJet_PRIME", "StarAir_GCR", "BatchSurcharges"} <= tables
