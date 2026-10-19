"""
The two extra rate sources the Flight Schedule prices with:

* Allcargo rate quotation (Index Global, .xlsx): one row per origin, destination,
  airline and flight (a flight number, ALL, or a product such as GCR). Its
  `Terminal&Surcharges` column is added to the airline's own tariff rate.
* Akasa special rates (.pdf): an SSP rate for +100 kg on a few Akasa routes,
  for Prime flights, GCR flights or ALL flights.
"""

import collections
import re

import openpyxl
import pdfplumber

# Schedule code for stations the rate sheets spell differently
STATION_ALIASES = {"TRV": "TVM"}


def airline_code(name):
    """'Air India' -> 'AIRINDIA', 'STAR AIR' -> 'STARAIR': the codes the schedule uses."""
    return re.sub(r"[^A-Z0-9]", "", str(name or "").upper())


def station(code):
    code = str(code or "").strip().upper()
    return STATION_ALIASES.get(code, code)


def flight_key(text):
    """'AI 1721' -> 'AI1721', '6E 0496' -> '6E496'; anything else (ALL, GCR, Prime) upper-cased."""
    s = str(text or "").strip().upper()
    m = re.fullmatch(r"([A-Z0-9]{2})[\s-]*0*(\d+)", s)
    return f"{m.group(1)}{m.group(2)}" if m else s


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r"\s*(\d+(?:\.\d+)?)", str(v or ""))
    return float(m.group(1)) if m else None


def _wef(texts):
    for t in texts:
        m = re.search(r"w\.?\s*e\.?\s*f\.?\s*([0-9]{1,2}[-./][0-9]{1,2}[-./][0-9]{2,4})", str(t or ""), re.I)
        if m:
            return m.group(1)
    return ""


def _header(row):
    return [re.sub(r"[^a-z0-9+&]", "", str(v or "").lower()) for v in row]


def _is_quote_header(row):
    h = _header(row)
    return "org" in h and "flight" in h and any("terminal" in x for x in h)


def is_quote(source):
    """True for an Allcargo-style quotation: an ORG / Flight table with a Terminal & Surcharges column."""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        return any(_is_quote_header(r) for ws in wb.worksheets for r in ws.iter_rows(max_row=40, values_only=True))
    finally:
        wb.close()


def read_quote(source):
    """{'wef': '01-10-2026', 'rows': {(org, dest, airline, flight): {'ts': float|None, 'note': str}}}"""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        ws = next(ws for ws in wb.worksheets
                  if any(_is_quote_header(r) for r in ws.iter_rows(max_row=40, values_only=True)))
        rows = ws.iter_rows(values_only=True)
        top = []
        for row in rows:
            h = _header(row)
            if _is_quote_header(row):
                col = {"org": h.index("org"), "dest": h.index("dest"), "airline": h.index("airline"),
                       "flight": h.index("flight"), "ts": next(i for i, x in enumerate(h) if "terminal" in x)}
                note = next((i for i, x in enumerate(h) if x.startswith("disckg")), None)
                break
            top.append(row[0])
        out = {}
        for row in rows:
            o, d, a = (row[col[k]] for k in ("org", "dest", "airline"))
            if not (isinstance(o, str) and len(o.strip()) == 3 and d and a):
                continue
            key = (station(o), station(d), airline_code(a), flight_key(row[col["flight"]]))
            ts = row[col["ts"]]
            out[key] = {"ts": float(ts) if isinstance(ts, (int, float)) else None,
                        "note": str(row[note]).strip() if note is not None and row[note] else ""}
        return {"wef": _wef(top), "rows": out, "parts": route_parts(out)}
    finally:
        wb.close()


def route_parts(rows):
    """Per airline, an origin part and a destination part such that a route's Terminal & Surcharges is
    origin part + destination part. That holds exactly on almost every ALL row of the quotation (e.g. every
    Akasa DEL value is 4.60 above the AYJ value to the same destination), so it gives the T&S for routes
    where the sheet has no number. -> {airline: ({origin: part}, {destination: part})}"""
    out = {}
    for airline in {k[2] for k in rows}:
        ts = {k[:2]: v["ts"] for k, v in rows.items() if k[2] == airline and k[3] == "ALL" and v["ts"] is not None}
        by_o = collections.defaultdict(dict)
        for (o, d), t in ts.items():
            by_o[o][d] = t
        if not by_o:
            continue
        ref = max(sorted(by_o), key=lambda o: len(by_o[o]))
        origin, todo = {ref: 0.0}, set(by_o) - {ref}
        while todo:  # place each origin by its usual difference to an origin already placed
            placed = False
            for o in sorted(todo):
                diffs = collections.Counter(round(by_o[o][d] - by_o[p][d] + origin[p], 2)
                                            for p in origin for d in by_o[o] if d in by_o[p])
                if diffs:
                    origin[o] = diffs.most_common(1)[0][0]
                    todo.discard(o)
                    placed = True
            if not placed:
                break
        dest = {}
        for d in {d for _, d in ts}:
            c = collections.Counter(round(t - origin[o], 2) for (o, dd), t in ts.items() if dd == d and o in origin)
            if c:
                dest[d] = c.most_common(1)[0][0]
        out[airline] = (origin, dest)
    return out


def read_special(source):
    """{'wef': ..., 'rows': {(org, dest, airline, 'PRIME'|'GCR'|'ALL'): {'ssp': float, 'console': float}}}"""
    with pdfplumber.open(source) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        tables = [t for p in pdf.pages for t in p.extract_tables()]
    out = {}
    for table in tables:
        for row in table:
            cells = [str(c or "").strip() for c in row]
            if len(cells) < 7 or len(cells[0]) != 3 or not cells[0].isalpha() or cells[0].upper() == "ORG":
                continue
            o, d, a, flight = cells[0], cells[1], cells[2], cells[3]
            ssp, console = _num(cells[5]), _num(cells[6])
            if ssp is None:
                continue
            out[(station(o), station(d), airline_code(a), flight_key(flight))] = {"ssp": ssp, "console": console}
    return {"wef": _wef([text]), "rows": out}


def terminal_surcharge(quote, flight, product):
    """(T&S ₹/kg, which row it came from) for a schedule flight: its own flight row, else its product's
    row (e.g. GCR), else the route's ALL row. Rows without a number are skipped. When the sheet has no
    number for the route, it is worked out from the airline's origin and destination parts ('DERIVED')."""
    if not quote:
        return None, None
    o, d, a = station(flight["Origin"]), station(flight["Destination"]), flight["Airline"]
    for k in (flight_key(flight["FlightNo"]), (product or "").upper(), "ALL"):
        row = quote["rows"].get((o, d, a, k))
        if row and row["ts"] is not None:
            return row["ts"], k
    origin, dest = quote.get("parts", {}).get(a, ({}, {}))
    if o in origin and d in dest:
        return round(origin[o] + dest[d], 2), "DERIVED"
    return None, None


def special_rate(special, flight, product):
    """The Akasa special rate for this flight, or None: the route's ALL row, else its product's row."""
    if not special:
        return None, None
    o, d, a = station(flight["Origin"]), station(flight["Destination"]), flight["Airline"]
    for k in ("ALL", (product or "").upper()):
        row = special["rows"].get((o, d, a, k))
        if row:
            return row, k
    return None, None
